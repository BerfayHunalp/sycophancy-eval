"""Build docs/Methodology_APA.docx from docs/methodology.md.

    py scripts/build_docx.py                 # pandoc -> docx, then APA-7 formatting pass
    py scripts/build_docx.py --src docs/other.md --out docs/Other.docx

Step 1: pandoc converts the Markdown (with raw ```{=openxml} page breaks) to docx.
Step 2: python-docx normalises the result to APA 7 professional-paper conventions:
Times New Roman 12 throughout, double spacing, 1-inch margins, running head + page
number in the header, centred bold Level-1 headings, left bold Level-2, bold-italic
Level-3, hanging indent on the reference list, single-spaced tables.
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "methodology.md"
OUT = ROOT / "docs" / "Methodology_APA.docx"
FONT = "Times New Roman"
SIZE = Pt(12)
RUNNING_HEAD = "PERSONA-INDUCED SYCOPHANCY IN FRONTIER LLMS"


def pandoc(src: Path, out: Path) -> None:
    exe = shutil.which("pandoc")
    if not exe:
        sys.exit("pandoc not found on PATH (winget install JohnMacFarlane.Pandoc)")
    cmd = [exe, str(src), "-f", "markdown+pipe_tables+raw_attribute", "-t", "docx", "-o", str(out),
           "--wrap=none"]
    subprocess.run(cmd, check=True)


def set_run_font(run, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT
    run.font.size = SIZE
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def style_paragraph(p, align=None, first_line: float | None = None, hanging: bool = False,
                    double: bool = True, space_after: int = 0) -> None:
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE if double else WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if hanging:
        pf.left_indent = Inches(0.5)
        pf.first_line_indent = Inches(-0.5)
    elif first_line is not None:
        pf.left_indent = Inches(0)
        pf.first_line_indent = Inches(first_line)


def add_page_number_header(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(RUNNING_HEAD)
        set_run_font(run)
        # right-aligned page number via a tab stop at the right margin
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), alignment=2)  # 2 = right
        run = p.add_run("\t")
        set_run_font(run)
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run = p.add_run()
        set_run_font(run)
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)


def format_body(doc: Document) -> None:
    in_refs = False
    for p in doc.paragraphs:
        name = (p.style.name or "").lower()
        if name.startswith("heading 1") or name == "title":
            for r in p.runs:
                set_run_font(r, bold=True)
            style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER)
            in_refs = p.text.strip().lower() in ("references", "reference list")
        elif name.startswith("heading 2"):
            for r in p.runs:
                set_run_font(r, bold=True)
            style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif name.startswith("heading 3"):
            for r in p.runs:
                set_run_font(r, bold=True, italic=True)
            style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT)
        elif name in ("author", "date", "subtitle"):
            for r in p.runs:
                set_run_font(r)
            style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            for r in p.runs:
                set_run_font(r)
            if in_refs and p.text.strip():
                style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, hanging=True)
            elif name in ("first paragraph", "body text", "normal", "compact") and p.text.strip():
                style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line=0.5)
            else:
                style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT)


def format_tables(doc: Document) -> None:
    for table in doc.tables:
        table.style = doc.styles["Table Grid"] if "Table Grid" in [s.name for s in doc.styles] else table.style
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_run_font(r)
                    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, double=False, space_after=0)
        # bold header row
        for cell in table.rows[0].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True


def set_default_style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=str(SRC))
    ap.add_argument("--out", default=str(OUT))
    args = ap.parse_args()
    src, out = Path(args.src), Path(args.out)
    if not src.exists():
        sys.exit(f"missing {src}")
    out.parent.mkdir(parents=True, exist_ok=True)
    pandoc(src, out)
    doc = Document(str(out))
    set_default_style(doc)
    add_page_number_header(doc)
    format_body(doc)
    format_tables(doc)
    doc.save(str(out))
    n_par = sum(1 for p in doc.paragraphs if p.text.strip())
    print(f"wrote {out.relative_to(ROOT)}: {n_par} paragraphs, {len(doc.tables)} tables, "
          f"{len(doc.sections)} section(s)")


if __name__ == "__main__":
    main()
